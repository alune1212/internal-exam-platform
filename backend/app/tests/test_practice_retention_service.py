import json
import zipfile
from datetime import UTC, datetime, timedelta
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.database import get_db
from app.core.security import create_admin_token
from app.models import PracticeAnswer, PracticeAnswerAggregate
from app.ops import internal_backup
from app.schemas.operations import MAX_RETENTION_SELECTION_IDS
from app.schemas.practice_retention import (
    PracticeRetentionArchiveRequest,
    PracticeRetentionDeleteRequest,
)
from app.services import backup_service, practice_retention_service
from app.tests.conftest import create_candidate, create_question_with_options


def _add_answer(
    db: Session,
    *,
    candidate_id: int,
    question_id: int,
    practiced_at: datetime,
    selected_answer: str = "A",
    is_correct: bool = False,
) -> PracticeAnswer:
    answer = PracticeAnswer(
        candidate_id=candidate_id,
        question_id=question_id,
        selected_answer=selected_answer,
        is_correct=is_correct,
        practiced_at=practiced_at,
    )
    db.add(answer)
    db.flush()
    return answer


def _add_aggregate(
    db: Session,
    *,
    candidate_id: int,
    question_id: int,
    answer: PracticeAnswer,
    total_attempts: int = 1,
    incorrect_count: int = 1,
) -> PracticeAnswerAggregate:
    aggregate = PracticeAnswerAggregate(
        candidate_id=candidate_id,
        question_id=question_id,
        total_attempts=total_attempts,
        incorrect_count=incorrect_count,
        latest_selected_answer=answer.selected_answer,
        latest_is_correct=answer.is_correct,
        latest_practiced_at=answer.practiced_at,
        latest_practice_answer_id=answer.id,
    )
    db.add(aggregate)
    return aggregate


def _verified_backup(root: Path, *, created_at: datetime) -> str:
    backup_id = "backup-practice-retention-test"
    directory = root / backup_id
    directory.mkdir(parents=True)
    (directory / internal_backup.DATABASE_DUMP_NAME).write_bytes(b"database")
    (directory / internal_backup.MEDIA_ARCHIVE_NAME).write_bytes(b"media")
    internal_backup.finalize_backup(
        directory,
        {
            "format_version": 1,
            "created_at": created_at.isoformat(),
            "migration_head": "202608300001",
            "table_counts": dict.fromkeys(internal_backup.TABLE_NAMES, 0),
            "media_file_count": 0,
        },
    )
    return backup_id


def test_practice_retention_preview_is_deterministic_and_reaches_low_watermark(
    db: Session,
) -> None:
    now = datetime(2026, 8, 30, 12, tzinfo=UTC)
    candidate = create_candidate(db)
    question = create_question_with_options(db)
    old = now - timedelta(days=366)
    recent = now - timedelta(days=100)
    answers = [
        _add_answer(
            db,
            candidate_id=candidate.id,
            question_id=question.id,
            practiced_at=old + timedelta(minutes=index),
        )
        for index in range(100)
    ]
    answers.extend(
        _add_answer(
            db,
            candidate_id=candidate.id,
            question_id=question.id,
            practiced_at=recent + timedelta(minutes=index),
        )
        for index in range(4900)
    )
    _add_aggregate(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        answer=answers[-1],
        total_attempts=len(answers),
        incorrect_count=len(answers),
    )
    db.commit()

    first = practice_retention_service.preview_practice_retention(db, now=now)
    second = practice_retention_service.preview_practice_retention(db, now=now)

    row = first.candidates[0]
    assert first.fingerprint == second.fingerprint
    assert row.hot_row_count == 5000
    assert row.expired_row_count == 100
    assert row.selected_row_count == 1000
    assert row.selected_practice_answer_ids == [answer.id for answer in answers[:1000]]


def test_practice_retention_preview_is_pageable_without_orm_history_materialization(
    db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 8, 30, 12, tzinfo=UTC)
    candidates = [create_candidate(db) for _ in range(3)]
    question = create_question_with_options(db)
    for candidate in candidates:
        answer = _add_answer(
            db,
            candidate_id=candidate.id,
            question_id=question.id,
            practiced_at=now - timedelta(days=366),
        )
        _add_aggregate(
            db,
            candidate_id=candidate.id,
            question_id=question.id,
            answer=answer,
        )
    db.commit()
    monkeypatch.setattr(
        db,
        "query",
        lambda *_args, **_kwargs: pytest.fail("preview must not ORM-load histories"),
    )

    preview = practice_retention_service.preview_practice_retention(
        db,
        now=now,
        limit=1,
        offset=1,
    )

    assert [row.candidate_id for row in preview.candidates] == [candidates[1].id]
    assert preview.candidates[0].selected_practice_answer_ids
    assert preview.limit == 1
    assert preview.offset == 1
    assert preview.total_candidates == 3
    assert preview.has_more is True


def test_practice_retention_preview_scope_ignores_unrelated_candidate_history(
    db: Session,
) -> None:
    now = datetime(2026, 8, 30, 12, tzinfo=UTC)
    selected_candidate = create_candidate(db)
    unrelated_candidate = create_candidate(db)
    question = create_question_with_options(db)
    selected_answer = _add_answer(
        db,
        candidate_id=selected_candidate.id,
        question_id=question.id,
        practiced_at=now - timedelta(days=366),
    )
    unrelated_answer = _add_answer(
        db,
        candidate_id=unrelated_candidate.id,
        question_id=question.id,
        practiced_at=now - timedelta(days=366),
    )
    _add_aggregate(
        db,
        candidate_id=selected_candidate.id,
        question_id=question.id,
        answer=selected_answer,
    )
    _add_aggregate(
        db,
        candidate_id=unrelated_candidate.id,
        question_id=question.id,
        answer=unrelated_answer,
    )
    db.commit()

    preview = practice_retention_service.preview_practice_retention(
        db,
        now=now,
        candidate_ids=[selected_candidate.id],
    )

    assert [row.candidate_id for row in preview.candidates] == [selected_candidate.id]
    assert preview.candidates[0].selected_practice_answer_ids == [selected_answer.id]
    assert preview.total_candidates == 1


def test_practice_retention_archive_delete_scope_selected_candidate_only(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 8, 30, 12, tzinfo=UTC)
    selected_candidate = create_candidate(db)
    unrelated_candidate = create_candidate(db)
    question = create_question_with_options(db)
    selected_answer = _add_answer(
        db,
        candidate_id=selected_candidate.id,
        question_id=question.id,
        practiced_at=now - timedelta(days=366),
    )
    selected_answer_id = selected_answer.id
    unrelated_answer = _add_answer(
        db,
        candidate_id=unrelated_candidate.id,
        question_id=question.id,
        practiced_at=now - timedelta(days=366),
    )
    _add_aggregate(
        db,
        candidate_id=selected_candidate.id,
        question_id=question.id,
        answer=selected_answer,
    )
    _add_aggregate(
        db,
        candidate_id=unrelated_candidate.id,
        question_id=question.id,
        answer=unrelated_answer,
    )
    db.commit()
    archive_root = tmp_path / "archives"
    backup_root = tmp_path / "backups"
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(archive_root))
    monkeypatch.setattr(settings, "backup_storage_dir", str(backup_root))
    preview = practice_retention_service.preview_practice_retention(
        db,
        now=now,
        candidate_ids=[selected_candidate.id],
    )
    archive = practice_retention_service.create_practice_retention_archive(
        db,
        candidate_ids=[selected_candidate.id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="admin",
        now=now,
    )
    backup_id = _verified_backup(backup_root, created_at=now + timedelta(minutes=1))

    deleted = practice_retention_service.delete_practice_retention(
        db,
        candidate_ids=[selected_candidate.id],
        preview_fingerprint=preview.fingerprint,
        archive_id=archive.artifact_id,
        backup_id=backup_id,
        confirmation=f"DELETE PRACTICE ANSWERS {selected_candidate.id}",
        operator_subject="admin",
        now=now + timedelta(minutes=2),
    )

    assert deleted.deleted_practice_answer_ids == [selected_answer_id]
    assert db.get(PracticeAnswer, selected_answer_id) is None
    assert db.get(PracticeAnswer, unrelated_answer.id) is not None


def test_practice_retention_archive_preserves_raw_json_and_escapes_workbook(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 8, 30, 12, tzinfo=UTC)
    candidate = create_candidate(db)
    question = create_question_with_options(db, stem="=SUM(A1:A2)")
    answer = _add_answer(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        practiced_at=now - timedelta(days=366),
        selected_answer="=1+1",
    )
    _add_aggregate(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        answer=answer,
    )
    db.commit()
    archive_root = tmp_path / "archives"
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(archive_root))

    preview = practice_retention_service.preview_practice_retention(db, now=now)
    archive = practice_retention_service.create_practice_retention_archive(
        db,
        candidate_ids=[candidate.id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="admin",
        now=now,
    )
    archive_path = archive_root / f"{archive.artifact_id}.zip"
    with zipfile.ZipFile(archive_path) as bundle:
        assert set(bundle.namelist()) == {
            "archive.json",
            "archive.xlsx",
            "manifest.json",
        }
        payload = json.loads(bundle.read("archive.json"))
        assert payload["practice_answers"][0]["selected_answer"] == "=1+1"
        workbook = load_workbook(bundle.open("archive.xlsx"))
        detail_sheet = workbook["练习明细"]
        assert detail_sheet.cell(row=2, column=4).value == "'=1+1"


def test_practice_retention_delete_is_fail_closed_then_preserves_aggregate(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 8, 30, 12, tzinfo=UTC)
    candidate = create_candidate(db)
    question = create_question_with_options(db)
    answer = _add_answer(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        practiced_at=now - timedelta(days=366),
    )
    answer_id = answer.id
    aggregate = _add_aggregate(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        answer=answer,
    )
    db.commit()
    archive_root = tmp_path / "archives"
    backup_root = tmp_path / "backups"
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(archive_root))
    monkeypatch.setattr(settings, "backup_storage_dir", str(backup_root))
    preview = practice_retention_service.preview_practice_retention(db, now=now)
    archive = practice_retention_service.create_practice_retention_archive(
        db,
        candidate_ids=[candidate.id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="admin",
        now=now,
    )

    answer.selected_answer = "tampered"
    db.commit()
    backup_id = _verified_backup(
        backup_root,
        created_at=now + timedelta(minutes=1),
    )
    with pytest.raises(practice_retention_service.PracticeRetentionSafeguardError):
        practice_retention_service.delete_practice_retention(
            db,
            candidate_ids=[candidate.id],
            preview_fingerprint=preview.fingerprint,
            archive_id=archive.artifact_id,
            backup_id=backup_id,
            confirmation=f"DELETE PRACTICE ANSWERS {candidate.id}",
            operator_subject="admin",
            now=now + timedelta(minutes=2),
        )
    assert db.get(PracticeAnswer, answer.id) is not None

    answer.selected_answer = "A"
    db.commit()
    deleted = practice_retention_service.delete_practice_retention(
        db,
        candidate_ids=[candidate.id],
        preview_fingerprint=preview.fingerprint,
        archive_id=archive.artifact_id,
        backup_id=backup_id,
        confirmation=f"DELETE PRACTICE ANSWERS {candidate.id}",
        operator_subject="admin",
        now=now + timedelta(minutes=2),
    )
    assert deleted.deleted_practice_answer_ids == [answer_id]
    assert db.get(PracticeAnswer, answer_id) is None
    assert db.get(PracticeAnswerAggregate, aggregate.id) is not None


def test_practice_retention_delete_requires_a_verified_newer_backup(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 8, 30, 12, tzinfo=UTC)
    candidate = create_candidate(db)
    question = create_question_with_options(db)
    answer = _add_answer(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        practiced_at=now - timedelta(days=366),
    )
    _add_aggregate(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        answer=answer,
    )
    db.commit()
    archive_root = tmp_path / "archives"
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(archive_root))
    monkeypatch.setattr(settings, "backup_storage_dir", str(tmp_path / "backups"))
    preview = practice_retention_service.preview_practice_retention(db, now=now)
    archive = practice_retention_service.create_practice_retention_archive(
        db,
        candidate_ids=[candidate.id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="admin",
        now=now,
    )

    with pytest.raises(practice_retention_service.PracticeRetentionSafeguardError):
        practice_retention_service.delete_practice_retention(
            db,
            candidate_ids=[candidate.id],
            preview_fingerprint=preview.fingerprint,
            archive_id=archive.artifact_id,
            backup_id="backup-missing",
            confirmation=f"DELETE PRACTICE ANSWERS {candidate.id}",
            operator_subject="admin",
            now=now + timedelta(minutes=2),
        )
    assert db.get(PracticeAnswer, answer.id) is not None


def test_practice_retention_delete_rejects_state_added_after_archive(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 8, 30, 12, tzinfo=UTC)
    candidate = create_candidate(db)
    question = create_question_with_options(db)
    answer = _add_answer(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        practiced_at=now - timedelta(days=366),
    )
    aggregate = _add_aggregate(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        answer=answer,
    )
    db.commit()
    archive_root = tmp_path / "archives"
    backup_root = tmp_path / "backups"
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(archive_root))
    monkeypatch.setattr(settings, "backup_storage_dir", str(backup_root))
    preview = practice_retention_service.preview_practice_retention(db, now=now)
    archive = practice_retention_service.create_practice_retention_archive(
        db,
        candidate_ids=[candidate.id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="admin",
        now=now,
    )
    new_answer = _add_answer(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        practiced_at=now - timedelta(minutes=1),
        is_correct=True,
    )
    aggregate.total_attempts += 1
    aggregate.latest_selected_answer = new_answer.selected_answer
    aggregate.latest_is_correct = new_answer.is_correct
    aggregate.latest_practiced_at = new_answer.practiced_at
    aggregate.latest_practice_answer_id = new_answer.id
    db.commit()
    _verified_backup(backup_root, created_at=now + timedelta(minutes=1))

    with pytest.raises(practice_retention_service.PracticeRetentionSafeguardError):
        practice_retention_service.delete_practice_retention(
            db,
            candidate_ids=[candidate.id],
            preview_fingerprint=preview.fingerprint,
            archive_id=archive.artifact_id,
            backup_id="backup-practice-retention-test",
            confirmation=f"DELETE PRACTICE ANSWERS {candidate.id}",
            operator_subject="admin",
            now=now + timedelta(minutes=2),
        )
    assert db.get(PracticeAnswer, answer.id) is not None
    assert db.get(PracticeAnswer, new_answer.id) is not None


def test_practice_retention_routes_require_admin_and_expose_preview(
    db: Session,
) -> None:
    from app.main import create_app

    app = create_app()

    def override_get_db():
        yield db

    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)
    assert (
        client.get("/api/admin/operations/practice-retention/preview").status_code
        == 401
    )
    token = create_admin_token(settings.admin_username)
    response = client.get(
        "/api/admin/operations/practice-retention/preview",
        headers={"X-Admin-Token": token},
    )
    assert response.status_code == 200
    assert response.json()["data"]["candidates"] == []

    candidate_ids_parameter = next(
        parameter
        for parameter in app.openapi()["paths"][
            "/api/admin/operations/practice-retention/preview"
        ]["get"]["parameters"]
        if parameter["name"] == "candidate_ids"
    )
    candidate_ids_schema = candidate_ids_parameter["schema"]["anyOf"][0]
    assert candidate_ids_schema["minItems"] == 1
    assert candidate_ids_schema["maxItems"] == MAX_RETENTION_SELECTION_IDS


def test_practice_retention_selection_limit_is_enforced_before_database_access(
    db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        practice_retention_service,
        "assert_admin_mutation_allowed",
        lambda _db: pytest.fail("selection limit must run before database access"),
    )
    oversized_ids = list(range(1, MAX_RETENTION_SELECTION_IDS + 2))

    with pytest.raises(
        practice_retention_service.PracticeRetentionSafeguardError, match="最多"
    ):
        practice_retention_service.preview_practice_retention(
            db,
            candidate_ids=oversized_ids,
        )

    with pytest.raises(ValidationError):
        PracticeRetentionArchiveRequest(
            candidate_ids=oversized_ids,
            preview_fingerprint="fingerprint",
        )
    with pytest.raises(ValidationError):
        PracticeRetentionDeleteRequest(
            candidate_ids=oversized_ids,
            preview_fingerprint="fingerprint",
            archive_id="practice-retention-20260830t120000z-aaaaaaaaaaaa",
            backup_id="backup-practice-retention-test",
            confirmation="DELETE PRACTICE ANSWERS 1",
        )

    with pytest.raises(
        practice_retention_service.PracticeRetentionSafeguardError, match="最多"
    ):
        practice_retention_service.create_practice_retention_archive(
            db,
            candidate_ids=oversized_ids,
            preview_fingerprint="fingerprint",
            operator_subject="admin",
        )
    with pytest.raises(
        practice_retention_service.PracticeRetentionSafeguardError, match="最多"
    ):
        practice_retention_service.delete_practice_retention(
            db,
            candidate_ids=oversized_ids,
            preview_fingerprint="fingerprint",
            archive_id="practice-retention-20260830t120000z-aaaaaaaaaaaa",
            backup_id="backup-practice-retention-test",
            confirmation="DELETE PRACTICE ANSWERS 1",
            operator_subject="admin",
        )


def test_practice_retention_selection_limit_accepts_exact_boundary() -> None:
    candidate_ids = list(range(1, MAX_RETENTION_SELECTION_IDS + 1))
    assert (
        practice_retention_service._normalize_candidate_ids(candidate_ids)
        == candidate_ids
    )
    assert (
        PracticeRetentionArchiveRequest(
            candidate_ids=candidate_ids, preview_fingerprint="fingerprint"
        ).candidate_ids
        == candidate_ids
    )
    assert (
        PracticeRetentionDeleteRequest(
            candidate_ids=candidate_ids,
            preview_fingerprint="fingerprint",
            archive_id="practice-retention-20260830t120000z-aaaaaaaaaaaa",
            backup_id="backup-practice-retention-test",
            confirmation="DELETE PRACTICE ANSWERS 1",
        ).candidate_ids
        == candidate_ids
    )


def test_practice_aggregate_changes_backup_data_fingerprint(
    db: Session, tmp_path: Path
) -> None:
    candidate = create_candidate(db)
    question = create_question_with_options(db)
    answer = _add_answer(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        practiced_at=datetime.now(UTC),
    )
    db.commit()
    before = backup_service.data_change_fingerprint(db, tmp_path / "media")
    aggregate = _add_aggregate(
        db,
        candidate_id=candidate.id,
        question_id=question.id,
        answer=answer,
    )
    db.commit()
    after = backup_service.data_change_fingerprint(db, tmp_path / "media")
    assert before != after
    assert "practice_answer_aggregate" in backup_service.DATA_TABLES

    aggregate.latest_practiced_at = datetime.now(UTC) + timedelta(seconds=1)
    db.commit()
    updated = backup_service.data_change_fingerprint(db, tmp_path / "media")
    assert updated != after
