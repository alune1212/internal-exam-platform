import hashlib
import json
import zipfile
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models import (
    Candidate,
    Exam,
    ExamAttempt,
    ExamAttemptAnswer,
    ExamAttemptQuestion,
    ExamCandidateScope,
)
from app.ops import internal_backup
from app.schemas.operations import (
    MAX_RETENTION_SELECTION_IDS,
    RetentionArchiveRequest,
    RetentionDeleteRequest,
)
from app.services import retention_service


def _old_exam_graph(db: Session, *, now: datetime) -> tuple[Exam, Candidate]:
    old = now - timedelta(days=400)
    candidate = Candidate(
        name="历史考试人", email="retention-candidate@example.com", status="active"
    )
    exam = Exam(
        title="历史考试",
        duration_minutes=60,
        status="archived",
        created_at=old,
        updated_at=old,
    )
    db.add_all([candidate, exam])
    db.flush()
    db.add(
        ExamCandidateScope(
            exam_id=exam.id,
            candidate_id=candidate.id,
            roster_email=candidate.email,
            roster_name=candidate.name or "待注册",
        )
    )
    attempt = ExamAttempt(
        exam_id=exam.id,
        candidate_id=candidate.id,
        status="submitted",
        started_at=old,
        ends_at=old + timedelta(hours=1),
        submitted_at=old + timedelta(hours=1),
        duration_minutes_snapshot=60,
        score=2,
        total_score=2,
        correct_count=1,
        wrong_count=0,
        created_at=old,
        updated_at=old,
    )
    db.add(attempt)
    db.flush()
    question = ExamAttemptQuestion(
        attempt_id=attempt.id,
        original_question_id=None,
        question_type="single",
        stem_snapshot="历史题干",
        options_snapshot=[{"label": "A", "content": "正确"}],
        correct_answer_snapshot="A",
        analysis_snapshot="历史解析",
        score=2,
        sort_order=1,
        created_at=old,
        updated_at=old,
    )
    db.add(question)
    db.flush()
    db.add(
        ExamAttemptAnswer(
            attempt_question_id=question.id,
            selected_answer="A",
            is_correct=True,
            score_awarded=2,
            answered_at=old,
            created_at=old,
            updated_at=old,
        )
    )
    db.commit()
    return exam, candidate


def _verified_backup(
    root: Path, *, created_at: datetime, backup_kind: str | None = None
) -> str:
    backup_id = "backup-retention-test"
    directory = root / backup_id
    directory.mkdir(parents=True)
    (directory / internal_backup.DATABASE_DUMP_NAME).write_bytes(b"database")
    (directory / internal_backup.MEDIA_ARCHIVE_NAME).write_bytes(b"media")
    manifest: dict[str, object] = {
        "format_version": 1,
        "created_at": created_at.isoformat(),
        "migration_head": "202607210001",
        "table_counts": dict.fromkeys(internal_backup.TABLE_NAMES, 0),
        "media_file_count": 0,
    }
    if backup_kind is not None:
        manifest.update(
            {
                "backup_kind": backup_kind,
                "dataset_id": "formal-dataset",
                "source_host_id": "source-host",
                "writer_generation": 4,
            }
        )
        if backup_kind == internal_backup.CUTOVER_BACKUP_KIND:
            manifest["writer_fence_boundary"] = {
                "dataset_id": "formal-dataset",
                "source_host_id": "source-host",
                "writer_generation": 4,
            }
    internal_backup.finalize_backup(directory, manifest)
    return backup_id


def test_retention_preview_archive_and_guarded_delete(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 7, 21, 8, tzinfo=UTC)
    exam, candidate = _old_exam_graph(db, now=now)
    exam_id = exam.id
    candidate_id = candidate.id
    scope_id = db.query(ExamCandidateScope).one().id
    archive_root = tmp_path / "archives"
    backup_root = tmp_path / "backups"
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(archive_root))
    monkeypatch.setattr(settings, "backup_storage_dir", str(backup_root))

    preview = retention_service.preview_retention(db, now=now)

    row = next(item for item in preview.exams if item.exam_id == exam_id)
    assert row.eligible is True
    assert row.attempt_count == 1
    assert row.answer_count == 1
    assert row.protected_candidate_count == 1

    archive = retention_service.create_retention_archive(
        db,
        exam_ids=[exam_id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="primary-operator",
        now=now,
    )
    archive_path = archive_root / f"{archive.artifact_id}.zip"
    with zipfile.ZipFile(archive_path) as bundle:
        assert set(bundle.namelist()) == {
            "archive.json",
            "archive.xlsx",
            "manifest.json",
        }
        exported = json.loads(bundle.read("archive.json"))
        assert exported["schema_version"] == 2
        assert exported["exams"][0]["scopes"][0] == {
            "scope_id": scope_id,
            "candidate_id": candidate_id,
            "roster_email": "retention-candidate@example.com",
            "roster_name": "历史考试人",
            "department": None,
            "position": None,
            "exam_group": None,
            "remark": None,
        }
        assert (
            exported["exams"][0]["attempts"][0]["questions"][0]["stem_snapshot"]
            == "历史题干"
        )
        workbook = load_workbook(BytesIO(bundle.read("archive.xlsx")), data_only=False)
        assert workbook["冻结名单"].max_row == 2
        assert (
            workbook["冻结名单"].cell(2, 4).value == "retention-candidate@example.com"
        )
        workbook.close()

    backup_id = _verified_backup(
        backup_root,
        created_at=now + timedelta(minutes=1),
        backup_kind=internal_backup.CUTOVER_BACKUP_KIND,
    )
    deleted = retention_service.delete_retained_exams(
        db,
        exam_ids=[exam_id],
        preview_fingerprint=preview.fingerprint,
        archive_id=archive.artifact_id,
        backup_id=backup_id,
        confirmation=f"DELETE EXAMS {exam_id}",
        operator_subject="primary-operator",
        now=now + timedelta(minutes=2),
    )

    assert deleted.deleted_exam_ids == [exam_id]
    assert deleted.deleted_attempt_count == 1
    assert deleted.protected_candidate_count == 1
    assert db.get(Exam, exam_id) is None
    assert db.get(Candidate, candidate_id) is not None


def test_retention_delete_fails_without_current_preview_archive_and_backup(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 7, 21, 8, tzinfo=UTC)
    exam, _ = _old_exam_graph(db, now=now)
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(tmp_path / "archives"))
    monkeypatch.setattr(settings, "backup_storage_dir", str(tmp_path / "backups"))
    preview = retention_service.preview_retention(db, now=now)

    with pytest.raises(retention_service.RetentionSafeguardError):
        retention_service.delete_retained_exams(
            db,
            exam_ids=[exam.id],
            preview_fingerprint="stale",
            archive_id="retention-20260721t080000z-aaaaaaaaaaaa",
            backup_id="missing",
            confirmation=f"DELETE EXAMS {exam.id}",
            operator_subject="primary-operator",
            now=now,
        )

    assert preview.exams[0].eligible is True
    assert db.get(Exam, exam.id) is not None


@pytest.mark.parametrize(
    "backup_id", ["../backup-outside", "/var/lib/internal-exam/backup-outside"]
)
def test_retention_delete_rejects_backup_path_before_validation(
    db: Session,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    backup_id: str,
) -> None:
    now = datetime(2026, 7, 21, 8, tzinfo=UTC)
    exam, _ = _old_exam_graph(db, now=now)
    archive_root = tmp_path / "archives"
    backup_root = tmp_path / "backups"
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(archive_root))
    monkeypatch.setattr(settings, "backup_storage_dir", str(backup_root))
    preview = retention_service.preview_retention(db, now=now)
    archive = retention_service.create_retention_archive(
        db,
        exam_ids=[exam.id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="primary-operator",
        now=now,
    )

    monkeypatch.setattr(
        retention_service,
        "validate_backup",
        lambda _path: pytest.fail("路径校验失败后不应读取配对备份"),
    )
    with pytest.raises(
        retention_service.RetentionSafeguardError, match="配对备份标识无效"
    ):
        retention_service.delete_retained_exams(
            db,
            exam_ids=[exam.id],
            preview_fingerprint=preview.fingerprint,
            archive_id=archive.artifact_id,
            backup_id=backup_id,
            confirmation=f"DELETE EXAMS {exam.id}",
            operator_subject="primary-operator",
            now=now + timedelta(minutes=2),
        )
    assert db.get(Exam, exam.id) is not None


def test_retention_delete_rejects_backup_symlink_before_validation(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 7, 21, 8, tzinfo=UTC)
    exam, _ = _old_exam_graph(db, now=now)
    archive_root = tmp_path / "archives"
    backup_root = tmp_path / "backups"
    backup_root.mkdir()
    outside_root = tmp_path / "outside"
    outside_root.mkdir()
    try:
        (backup_root / "backup-escaped").symlink_to(
            outside_root, target_is_directory=True
        )
    except OSError as exc:
        pytest.skip(f"当前环境不支持 symlink：{exc}")

    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(archive_root))
    monkeypatch.setattr(settings, "backup_storage_dir", str(backup_root))
    preview = retention_service.preview_retention(db, now=now)
    archive = retention_service.create_retention_archive(
        db,
        exam_ids=[exam.id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="primary-operator",
        now=now,
    )
    monkeypatch.setattr(
        retention_service,
        "validate_backup",
        lambda _path: pytest.fail("symlink 路径校验失败后不应读取配对备份"),
    )

    with pytest.raises(
        retention_service.RetentionSafeguardError, match="配对备份标识无效"
    ):
        retention_service.delete_retained_exams(
            db,
            exam_ids=[exam.id],
            preview_fingerprint=preview.fingerprint,
            archive_id=archive.artifact_id,
            backup_id="backup-escaped",
            confirmation=f"DELETE EXAMS {exam.id}",
            operator_subject="primary-operator",
            now=now + timedelta(minutes=2),
        )
    assert db.get(Exam, exam.id) is not None


def test_retention_archive_escapes_workbook_strings_but_keeps_json_raw(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 7, 21, 8, tzinfo=UTC)
    exam, _ = _old_exam_graph(db, now=now)
    scope = db.query(ExamCandidateScope).one()
    question = db.query(ExamAttemptQuestion).one()
    answer = db.query(ExamAttemptAnswer).one()
    exam.title = "\t=TITLE()"
    scope.roster_name = "\t=NAME()"
    scope.department = "+DEPARTMENT()"
    scope.position = "-POSITION()"
    scope.exam_group = "@GROUP()"
    scope.roster_remark = "=REMARK()"
    question.stem_snapshot = "\t=STEM()"
    question.correct_answer_snapshot = "+CORRECT()"
    answer.selected_answer = "-SELECTED()"
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(tmp_path / "archives"))

    preview = retention_service.preview_retention(db, now=now)
    archive = retention_service.create_retention_archive(
        db,
        exam_ids=[exam.id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="primary-operator",
        now=now,
    )
    with zipfile.ZipFile(
        tmp_path / "archives" / f"{archive.artifact_id}.zip"
    ) as bundle:
        exported = json.loads(bundle.read("archive.json"))
        assert exported["exams"][0]["title"] == "\t=TITLE()"
        assert exported["exams"][0]["scopes"][0]["roster_name"] == "\t=NAME()"
        workbook = load_workbook(BytesIO(bundle.read("archive.xlsx")), data_only=False)

    assert workbook["考试归档"].cell(2, 2).value == "'\t=TITLE()"
    assert workbook["冻结名单"].cell(2, 5).value == "'\t=NAME()"
    assert workbook["冻结名单"].cell(2, 6).value == "'+DEPARTMENT()"
    assert workbook["冻结名单"].cell(2, 7).value == "'-POSITION()"
    assert workbook["冻结名单"].cell(2, 8).value == "'@GROUP()"
    assert workbook["冻结名单"].cell(2, 9).value == "'=REMARK()"
    assert workbook["作答快照"].cell(2, 3).value == "'\t=STEM()"
    assert workbook["作答快照"].cell(2, 4).value == "'-SELECTED()"
    assert workbook["作答快照"].cell(2, 5).value == "'+CORRECT()"
    workbook.close()


def test_retention_delete_rejects_changed_current_roster(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 7, 21, 8, tzinfo=UTC)
    exam, _ = _old_exam_graph(db, now=now)
    archive_root = tmp_path / "archives"
    backup_root = tmp_path / "backups"
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(archive_root))
    monkeypatch.setattr(settings, "backup_storage_dir", str(backup_root))
    preview = retention_service.preview_retention(db, now=now)
    archive = retention_service.create_retention_archive(
        db,
        exam_ids=[exam.id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="primary-operator",
        now=now,
    )
    scope = db.query(ExamCandidateScope).one()
    scope.roster_name = "名单已变化"
    db.commit()
    backup_id = _verified_backup(
        backup_root,
        created_at=now + timedelta(minutes=1),
        backup_kind=internal_backup.CUTOVER_BACKUP_KIND,
    )

    with pytest.raises(retention_service.RetentionSafeguardError, match="来源"):
        retention_service.delete_retained_exams(
            db,
            exam_ids=[exam.id],
            preview_fingerprint=preview.fingerprint,
            archive_id=archive.artifact_id,
            backup_id=backup_id,
            confirmation=f"DELETE EXAMS {exam.id}",
            operator_subject="primary-operator",
            now=now + timedelta(minutes=2),
        )
    assert db.get(Exam, exam.id) is not None


def test_retention_schema_v1_is_readable_but_cannot_delete(
    db: Session, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    now = datetime(2026, 7, 21, 8, tzinfo=UTC)
    exam, _ = _old_exam_graph(db, now=now)
    archive_root = tmp_path / "archives"
    backup_root = tmp_path / "backups"
    monkeypatch.setattr(settings, "lifecycle_archive_dir", str(archive_root))
    monkeypatch.setattr(settings, "backup_storage_dir", str(backup_root))
    preview = retention_service.preview_retention(db, now=now)
    archive = retention_service.create_retention_archive(
        db,
        exam_ids=[exam.id],
        preview_fingerprint=preview.fingerprint,
        operator_subject="primary-operator",
        now=now,
    )
    archive_path = archive_root / f"{archive.artifact_id}.zip"
    with zipfile.ZipFile(archive_path) as bundle:
        members = {name: bundle.read(name) for name in bundle.namelist()}
    manifest = json.loads(members["manifest.json"])
    manifest["schema_version"] = 1
    manifest_bytes = retention_service._json_bytes(manifest)
    rewritten = BytesIO()
    with zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as bundle:
        bundle.writestr("archive.json", members["archive.json"])
        bundle.writestr("archive.xlsx", members["archive.xlsx"])
        bundle.writestr("manifest.json", manifest_bytes)
    archive_path.write_bytes(rewritten.getvalue())
    (archive_root / f"{archive.artifact_id}.manifest.json").write_bytes(manifest_bytes)
    (archive_root / f"{archive.artifact_id}.sha256").write_text(
        f"{hashlib.sha256(archive_path.read_bytes()).hexdigest()}  {archive_path.name}\n",
        encoding="ascii",
    )

    assert (
        retention_service._load_archive_manifest(archive.artifact_id)["schema_version"]
        == 1
    )
    backup_id = _verified_backup(
        backup_root,
        created_at=now + timedelta(minutes=1),
        backup_kind=internal_backup.CUTOVER_BACKUP_KIND,
    )
    with pytest.raises(retention_service.RetentionSafeguardError, match="仅可读取"):
        retention_service.delete_retained_exams(
            db,
            exam_ids=[exam.id],
            preview_fingerprint=preview.fingerprint,
            archive_id=archive.artifact_id,
            backup_id=backup_id,
            confirmation=f"DELETE EXAMS {exam.id}",
            operator_subject="primary-operator",
            now=now + timedelta(minutes=2),
        )
    assert db.get(Exam, exam.id) is not None


def test_retention_preview_explains_active_and_recent_exclusions(db: Session) -> None:
    now = datetime(2026, 7, 21, 8, tzinfo=UTC)
    db.add(
        Exam(
            title="近期正式考试",
            duration_minutes=60,
            status="active",
            created_at=now,
            updated_at=now,
        )
    )
    db.commit()

    preview = retention_service.preview_retention(db, now=now)

    assert preview.exams[0].eligible is False
    assert "考试尚未归档" in preview.exams[0].reasons
    assert "最终活动距今未满 12 个月" in preview.exams[0].reasons


def test_retention_selection_limit_is_enforced_before_database_access(
    db: Session, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        retention_service,
        "assert_admin_mutation_allowed",
        lambda _db: pytest.fail("selection limit must run before database access"),
    )
    oversized_ids = list(range(1, MAX_RETENTION_SELECTION_IDS + 2))

    with pytest.raises(ValidationError):
        RetentionArchiveRequest(
            exam_ids=oversized_ids,
            preview_fingerprint="fingerprint",
        )
    with pytest.raises(ValidationError):
        RetentionDeleteRequest(
            exam_ids=oversized_ids,
            preview_fingerprint="fingerprint",
            archive_id="retention-20260721t080000z-aaaaaaaaaaaa",
            backup_id="backup-retention-test",
            confirmation="DELETE EXAMS 1",
        )

    with pytest.raises(retention_service.RetentionSafeguardError, match="最多"):
        retention_service.create_retention_archive(
            db,
            exam_ids=oversized_ids,
            preview_fingerprint="fingerprint",
            operator_subject="primary-operator",
        )
    with pytest.raises(retention_service.RetentionSafeguardError, match="最多"):
        retention_service.delete_retained_exams(
            db,
            exam_ids=oversized_ids,
            preview_fingerprint="fingerprint",
            archive_id="retention-20260721t080000z-aaaaaaaaaaaa",
            backup_id="backup-retention-test",
            confirmation="DELETE EXAMS 1",
            operator_subject="primary-operator",
        )


def test_retention_selection_limit_accepts_exact_boundary() -> None:
    exam_ids = list(range(1, MAX_RETENTION_SELECTION_IDS + 1))
    assert retention_service._normalize_exam_ids(exam_ids) == exam_ids
    assert (
        RetentionArchiveRequest(
            exam_ids=exam_ids, preview_fingerprint="fingerprint"
        ).exam_ids
        == exam_ids
    )
    assert (
        RetentionDeleteRequest(
            exam_ids=exam_ids,
            preview_fingerprint="fingerprint",
            archive_id="retention-20260721t080000z-aaaaaaaaaaaa",
            backup_id="backup-retention-test",
            confirmation="DELETE EXAMS 1",
        ).exam_ids
        == exam_ids
    )
