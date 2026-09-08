from __future__ import annotations

import hashlib
import json
import re
import stat
import zipfile
from datetime import UTC, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from typing import TYPE_CHECKING, Any

from openpyxl import Workbook
from sqlalchemy import case, func, or_, select
from sqlalchemy import inspect as sqlalchemy_inspect

from app.core.config import settings
from app.core.exceptions import DomainError
from app.core.time import to_utc
from app.models import Candidate, PracticeAnswer, PracticeAnswerAggregate
from app.ops.internal_backup import BackupValidationError, validate_backup
from app.schemas.practice_retention import (
    MAX_RETENTION_SELECTION_IDS,
    PracticeRetentionArchiveRead,
    PracticeRetentionCandidatePreview,
    PracticeRetentionDeleteRead,
    PracticeRetentionPreviewRead,
)
from app.services.audit_service import record_admin_event
from app.services.excel_security import escape_excel_cell
from app.services.operational_lock_service import assert_admin_mutation_allowed

if TYPE_CHECKING:
    from sqlalchemy.orm import Session

RETENTION_DAYS = 365
HOT_HISTORY_LIMIT = 5000
HOT_HISTORY_LOW_WATERMARK = 4000
ARTIFACT_ID_PATTERN = re.compile(
    r"^practice-retention-[0-9]{8}t[0-9]{6}z-[0-9a-f]{12}$"
)
BACKUP_ID_PATTERN = re.compile(r"^backup-[A-Za-z0-9][A-Za-z0-9._:-]{0,127}$")
ARCHIVE_MEMBERS = frozenset({"archive.json", "archive.xlsx", "manifest.json"})
ARCHIVE_PAYLOAD_KEYS = frozenset(
    {
        "schema_version",
        "kind",
        "candidate_ids",
        "practice_answer_ids",
        "preview_fingerprint",
        "cutoff_at",
        "retention_days",
        "hot_history_limit",
        "hot_history_low_watermark",
        "practice_answers",
        "aggregate_snapshots",
    }
)
ARCHIVE_MANIFEST_KEYS = frozenset(
    {
        "schema_version",
        "kind",
        "artifact_id",
        "created_at",
        "operator_subject",
        "candidate_ids",
        "practice_answer_ids",
        "preview_fingerprint",
        "cutoff_at",
        "retention_days",
        "hot_history_limit",
        "hot_history_low_watermark",
        "files",
    }
)


class PracticeRetentionSafeguardError(DomainError):
    status_code = 409


def _json_value(value: object) -> object:
    if isinstance(value, datetime):
        return to_utc(value).isoformat()
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, dict):
        return {str(key): _json_value(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_json_value(item) for item in value]
    if isinstance(value, tuple):
        return [_json_value(item) for item in value]
    return value


def _json_bytes(payload: object) -> bytes:
    return (
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True, default=str)
        + "\n"
    ).encode("utf-8")


def _answer_snapshot(answer: PracticeAnswer) -> dict[str, object]:
    return {
        "id": answer.id,
        "candidate_id": answer.candidate_id,
        "question_id": answer.question_id,
        "selected_answer": answer.selected_answer,
        "is_correct": answer.is_correct,
        "practiced_at": to_utc(answer.practiced_at).isoformat(),
    }


def _aggregate_snapshot(aggregate: PracticeAnswerAggregate) -> dict[str, object]:
    mapper = sqlalchemy_inspect(type(aggregate)).mapper
    return {
        attribute.key: _json_value(getattr(aggregate, attribute.key))
        for attribute in mapper.column_attrs
    }


def _fingerprint(
    cutoff_at: datetime, candidates: list[PracticeRetentionCandidatePreview]
) -> str:
    # Keep a preview valid for the day it was generated; the selected IDs are
    # compared again during archive/delete if the exact 365-day boundary moves.
    fingerprint_cutoff = datetime.combine(cutoff_at.date(), time.min, UTC)
    payload = {
        "cutoff_at": fingerprint_cutoff.isoformat(),
        "retention_days": RETENTION_DAYS,
        "hot_history_limit": HOT_HISTORY_LIMIT,
        "hot_history_low_watermark": HOT_HISTORY_LOW_WATERMARK,
        "candidates": [candidate.model_dump(mode="json") for candidate in candidates],
    }
    return hashlib.sha256(
        json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
    ).hexdigest()


def _normalize_candidate_ids(candidate_ids: list[int]) -> list[int]:
    if len(candidate_ids) > MAX_RETENTION_SELECTION_IDS:
        raise PracticeRetentionSafeguardError(
            f"一次最多选择 {MAX_RETENTION_SELECTION_IDS} 个考试人。"
        )
    if not candidate_ids or any(
        isinstance(candidate_id, bool) or candidate_id < 1
        for candidate_id in candidate_ids
    ):
        raise PracticeRetentionSafeguardError("必须提供非空、唯一且有序的考试人 ID。")
    normalized = sorted(set(candidate_ids))
    if normalized != candidate_ids:
        raise PracticeRetentionSafeguardError("必须提供非空、唯一且有序的考试人 ID。")
    return normalized


def _candidate_count_rows(
    db: Session,
    *,
    cutoff_at: datetime,
    candidate_ids: list[int] | None,
    limit: int | None,
    offset: int,
) -> tuple[list[tuple[int, int, int]], int]:
    expired_case = case((PracticeAnswer.practiced_at < cutoff_at, 1), else_=0)
    grouped = select(
        PracticeAnswer.candidate_id.label("candidate_id"),
        func.count(PracticeAnswer.id).label("hot_row_count"),
        func.sum(expired_case).label("expired_row_count"),
    ).group_by(PracticeAnswer.candidate_id)
    if candidate_ids is not None:
        grouped = grouped.where(PracticeAnswer.candidate_id.in_(candidate_ids))

    total = int(
        db.execute(
            select(func.count()).select_from(grouped.order_by(None).subquery())
        ).scalar_one()
    )
    page = grouped.order_by(PracticeAnswer.candidate_id)
    if limit is not None:
        page = page.offset(offset).limit(limit)
    rows = [
        (
            int(row.candidate_id),
            int(row.hot_row_count),
            int(row.expired_row_count or 0),
        )
        for row in db.execute(page).all()
    ]
    return rows, total


def _selected_answer_ids(
    db: Session,
    *,
    candidate_ids: list[int],
    cutoff_at: datetime,
) -> dict[int, list[int]]:
    """Select retention rows in SQL without materializing unrelated details."""

    row_number = func.row_number().over(
        partition_by=PracticeAnswer.candidate_id,
        order_by=(PracticeAnswer.practiced_at.asc(), PracticeAnswer.id.asc()),
    )
    hot_count = func.count(PracticeAnswer.id).over(
        partition_by=PracticeAnswer.candidate_id
    )
    ranked = (
        select(
            PracticeAnswer.id.label("practice_answer_id"),
            PracticeAnswer.candidate_id.label("candidate_id"),
            PracticeAnswer.practiced_at.label("practiced_at"),
            row_number.label("row_number"),
            hot_count.label("hot_row_count"),
        )
        .where(PracticeAnswer.candidate_id.in_(candidate_ids))
        .subquery()
    )
    selection = (
        select(
            ranked.c.candidate_id,
            ranked.c.practice_answer_id,
        )
        .where(
            or_(
                ranked.c.practiced_at < cutoff_at,
                (
                    (ranked.c.hot_row_count >= HOT_HISTORY_LIMIT)
                    & (
                        ranked.c.row_number
                        <= ranked.c.hot_row_count - HOT_HISTORY_LOW_WATERMARK
                    )
                ),
            )
        )
        .order_by(
            ranked.c.candidate_id,
            ranked.c.practiced_at,
            ranked.c.practice_answer_id,
        )
    )
    selected: dict[int, list[int]] = {}
    for row in db.execute(selection).all():
        selected.setdefault(int(row.candidate_id), []).append(
            int(row.practice_answer_id)
        )
    return selected


def _preview_candidates(
    db: Session,
    *,
    cutoff_at: datetime,
    candidate_ids: list[int] | None,
    limit: int | None,
    offset: int,
) -> tuple[list[PracticeRetentionCandidatePreview], int]:
    count_rows, total = _candidate_count_rows(
        db,
        cutoff_at=cutoff_at,
        candidate_ids=candidate_ids,
        limit=limit,
        offset=offset,
    )
    selected = _selected_answer_ids(
        db,
        candidate_ids=[row[0] for row in count_rows],
        cutoff_at=cutoff_at,
    )
    candidates = [
        PracticeRetentionCandidatePreview(
            candidate_id=candidate_id,
            hot_row_count=hot_row_count,
            expired_row_count=expired_row_count,
            selected_row_count=len(selected.get(candidate_id, [])),
            selected_practice_answer_ids=selected.get(candidate_id, []),
            eligible=bool(selected.get(candidate_id)),
            capacity_limited=hot_row_count >= HOT_HISTORY_LIMIT,
        )
        for candidate_id, hot_row_count, expired_row_count in count_rows
    ]
    return candidates, total


def preview_practice_retention(
    db: Session,
    *,
    now: datetime | None = None,
    candidate_ids: list[int] | None = None,
    limit: int = 100,
    offset: int = 0,
) -> PracticeRetentionPreviewRead:
    generated_at = to_utc(now or datetime.now(UTC))
    cutoff_at = generated_at - timedelta(days=RETENTION_DAYS)
    if isinstance(limit, bool) or not isinstance(limit, int) or not 1 <= limit <= 100:
        raise PracticeRetentionSafeguardError("预览 limit 必须在 1 到 100 之间。")
    if isinstance(offset, bool) or not isinstance(offset, int) or offset < 0:
        raise PracticeRetentionSafeguardError("预览 offset 必须是非负整数。")
    normalized_ids = (
        _normalize_candidate_ids(candidate_ids) if candidate_ids is not None else None
    )
    candidates, total = _preview_candidates(
        db,
        cutoff_at=cutoff_at,
        candidate_ids=normalized_ids,
        limit=None if normalized_ids is not None else limit,
        offset=0 if normalized_ids is not None else offset,
    )
    return PracticeRetentionPreviewRead(
        generated_at=generated_at,
        cutoff_at=cutoff_at,
        retention_days=RETENTION_DAYS,
        hot_history_limit=HOT_HISTORY_LIMIT,
        hot_history_low_watermark=HOT_HISTORY_LOW_WATERMARK,
        fingerprint=_fingerprint(cutoff_at, candidates),
        candidates=candidates,
        limit=len(candidates) if normalized_ids is not None else limit,
        offset=0 if normalized_ids is not None else offset,
        total_candidates=total,
        has_more=(
            False if normalized_ids is not None else offset + len(candidates) < total
        ),
    )


def _selected_ids_for_candidates(
    preview: PracticeRetentionPreviewRead, candidate_ids: list[int]
) -> list[int]:
    rows_by_id = {row.candidate_id: row for row in preview.candidates}
    if any(
        candidate_id not in rows_by_id or not rows_by_id[candidate_id].eligible
        for candidate_id in candidate_ids
    ):
        raise PracticeRetentionSafeguardError("所选考试人没有符合条件的练习明细。")
    answer_ids = [
        answer_id
        for candidate_id in candidate_ids
        for answer_id in rows_by_id[candidate_id].selected_practice_answer_ids
    ]
    if len(answer_ids) != len(set(answer_ids)):
        raise PracticeRetentionSafeguardError("练习归档选择包含重复明细。")
    return answer_ids


def _load_archive_payload(
    db: Session,
    *,
    candidate_ids: list[int],
    answer_ids: list[int],
    preview: PracticeRetentionPreviewRead,
) -> dict[str, Any]:
    answers = (
        db.query(PracticeAnswer)
        .filter(
            PracticeAnswer.id.in_(answer_ids),
            PracticeAnswer.candidate_id.in_(candidate_ids),
        )
        .order_by(
            PracticeAnswer.candidate_id, PracticeAnswer.practiced_at, PracticeAnswer.id
        )
        .all()
    )
    if [answer.id for answer in answers] != answer_ids:
        raise PracticeRetentionSafeguardError("练习明细已发生变化，请重新预览。")

    pairs = {(answer.candidate_id, answer.question_id) for answer in answers}
    aggregate_rows = (
        db.query(PracticeAnswerAggregate)
        .filter(PracticeAnswerAggregate.candidate_id.in_(candidate_ids))
        .order_by(
            PracticeAnswerAggregate.candidate_id,
            PracticeAnswerAggregate.question_id,
        )
        .all()
    )
    aggregates = {
        (aggregate.candidate_id, aggregate.question_id): aggregate
        for aggregate in aggregate_rows
        if (aggregate.candidate_id, aggregate.question_id) in pairs
    }
    if set(aggregates) != pairs:
        raise PracticeRetentionSafeguardError("练习聚合状态不完整，拒绝创建归档。")

    return {
        "schema_version": 1,
        "kind": "practice-retention",
        "candidate_ids": candidate_ids,
        "practice_answer_ids": answer_ids,
        "preview_fingerprint": preview.fingerprint,
        "cutoff_at": preview.cutoff_at.isoformat(),
        "retention_days": RETENTION_DAYS,
        "hot_history_limit": HOT_HISTORY_LIMIT,
        "hot_history_low_watermark": HOT_HISTORY_LOW_WATERMARK,
        "practice_answers": [_answer_snapshot(answer) for answer in answers],
        "aggregate_snapshots": [
            _aggregate_snapshot(aggregates[pair]) for pair in sorted(aggregates)
        ],
    }


def _workbook_bytes(
    payload: dict[str, Any], preview: PracticeRetentionPreviewRead
) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "练习摘要"
    summary.append(["考试人ID", "热明细数", "过期明细数", "归档明细数", "达到容量上限"])
    for candidate in preview.candidates:
        if candidate.candidate_id not in payload["candidate_ids"]:
            continue
        summary.append(
            [
                candidate.candidate_id,
                candidate.hot_row_count,
                candidate.expired_row_count,
                candidate.selected_row_count,
                candidate.capacity_limited,
            ]
        )

    details = workbook.create_sheet("练习明细")
    details.append(["明细ID", "考试人ID", "题目ID", "所选答案", "是否正确", "练习时间"])
    for answer in payload["practice_answers"]:
        details.append(
            [
                answer["id"],
                answer["candidate_id"],
                answer["question_id"],
                answer["selected_answer"],
                answer["is_correct"],
                answer["practiced_at"],
            ]
        )

    aggregates = workbook.create_sheet("聚合快照")
    aggregate_rows = payload["aggregate_snapshots"]
    aggregate_keys = sorted(
        {key for row in aggregate_rows for key in row},
        key=str,
    )
    if aggregate_keys:
        aggregates.append(aggregate_keys)
        for row in aggregate_rows:
            aggregates.append([row.get(key) for key in aggregate_keys])

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.value = escape_excel_cell(cell.value)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def create_practice_retention_archive(
    db: Session,
    *,
    candidate_ids: list[int],
    preview_fingerprint: str,
    operator_subject: str,
    now: datetime | None = None,
) -> PracticeRetentionArchiveRead:
    normalized_ids = _normalize_candidate_ids(candidate_ids)
    assert_admin_mutation_allowed(db)
    try:
        created_at = to_utc(now or datetime.now(UTC))
        preview = preview_practice_retention(
            db,
            now=created_at,
            candidate_ids=normalized_ids,
        )
        if preview.fingerprint != preview_fingerprint:
            raise PracticeRetentionSafeguardError("保留预览已过期，请重新预览。")
        answer_ids = _selected_ids_for_candidates(preview, normalized_ids)
        payload = _load_archive_payload(
            db,
            candidate_ids=normalized_ids,
            answer_ids=answer_ids,
            preview=preview,
        )
        json_content = _json_bytes(payload)
        workbook_content = _workbook_bytes(payload, preview)
        artifact_id = (
            f"practice-retention-{created_at.strftime('%Y%m%dt%H%M%Sz').lower()}-"
            f"{preview.fingerprint[:12]}"
        )
        manifest = {
            "schema_version": 1,
            "kind": "practice-retention",
            "artifact_id": artifact_id,
            "created_at": created_at.isoformat(),
            "operator_subject": operator_subject,
            "candidate_ids": normalized_ids,
            "practice_answer_ids": answer_ids,
            "preview_fingerprint": preview.fingerprint,
            "cutoff_at": preview.cutoff_at.isoformat(),
            "retention_days": RETENTION_DAYS,
            "hot_history_limit": HOT_HISTORY_LIMIT,
            "hot_history_low_watermark": HOT_HISTORY_LOW_WATERMARK,
            "files": {
                "archive.json": hashlib.sha256(json_content).hexdigest(),
                "archive.xlsx": hashlib.sha256(workbook_content).hexdigest(),
            },
        }
        manifest_content = _json_bytes(manifest)
        archive_dir = Path(settings.lifecycle_archive_dir)
        archive_dir.mkdir(parents=True, exist_ok=True)
        archive_path = archive_dir / f"{artifact_id}.zip"
        temporary_path = archive_dir / f".{artifact_id}.tmp"
        with zipfile.ZipFile(
            temporary_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as bundle:
            bundle.writestr("archive.json", json_content)
            bundle.writestr("archive.xlsx", workbook_content)
            bundle.writestr("manifest.json", manifest_content)
        temporary_path.replace(archive_path)
        archive_sha256 = hashlib.sha256(archive_path.read_bytes()).hexdigest()
        (archive_dir / f"{artifact_id}.manifest.json").write_bytes(manifest_content)
        (archive_dir / f"{artifact_id}.sha256").write_text(
            f"{archive_sha256}  {archive_path.name}\n", encoding="ascii"
        )
        record_admin_event(
            db,
            operator_subject=operator_subject,
            action="practice_retention_archive_created",
            target_type="practice_answer_set",
            target_id=",".join(map(str, normalized_ids)),
            metadata={
                "archive_ref": artifact_id,
                "count": len(answer_ids),
                "selected_count": len(normalized_ids),
                "fingerprint": preview.fingerprint,
            },
        )
        db.commit()
        return PracticeRetentionArchiveRead(
            artifact_id=artifact_id,
            created_at=created_at,
            candidate_ids=normalized_ids,
            practice_answer_ids=answer_ids,
            preview_fingerprint=preview.fingerprint,
            archive_sha256=archive_sha256,
        )
    except PracticeRetentionSafeguardError:
        db.rollback()
        raise
    except Exception:
        db.rollback()
        raise


def _regular_file(path: Path) -> bool:
    try:
        return stat.S_ISREG(path.lstat().st_mode)
    except OSError:
        return False


def _load_archive(artifact_id: str) -> tuple[dict[str, Any], dict[str, Any], str]:
    if ARTIFACT_ID_PATTERN.fullmatch(artifact_id) is None:
        raise PracticeRetentionSafeguardError("练习归档产物标识无效。")
    archive_dir = Path(settings.lifecycle_archive_dir)
    archive_path = archive_dir / f"{artifact_id}.zip"
    manifest_path = archive_dir / f"{artifact_id}.manifest.json"
    checksum_path = archive_dir / f"{artifact_id}.sha256"
    if not all(
        _regular_file(path) for path in (archive_path, manifest_path, checksum_path)
    ):
        raise PracticeRetentionSafeguardError("练习归档产物不完整。")
    try:
        checksum_rows = checksum_path.read_text(encoding="ascii").splitlines()
        if len(checksum_rows) != 1:
            raise ValueError("invalid archive checksum")
        expected, filename = checksum_rows[0].split("  ", 1)
        if filename != archive_path.name or not re.fullmatch(r"[0-9a-f]{64}", expected):
            raise ValueError("invalid archive checksum")
    except (OSError, UnicodeError, ValueError) as exc:
        raise PracticeRetentionSafeguardError(
            "练习归档 checksum 文件格式无效。"
        ) from exc
    try:
        archive_sha256 = hashlib.sha256(archive_path.read_bytes()).hexdigest()
    except OSError as exc:
        raise PracticeRetentionSafeguardError("练习归档产物无法读取。") from exc
    if archive_sha256 != expected:
        raise PracticeRetentionSafeguardError("练习归档 checksum 校验失败。")

    try:
        outer_manifest_bytes = manifest_path.read_bytes()
        manifest = json.loads(outer_manifest_bytes)
        with zipfile.ZipFile(archive_path) as bundle:
            names = bundle.namelist()
            if len(names) != len(set(names)) or set(names) != ARCHIVE_MEMBERS:
                raise ValueError("invalid archive members")
            infos = {info.filename: info for info in bundle.infolist()}
            if any(
                info.is_dir() or stat.S_ISLNK((info.external_attr >> 16) & 0xFFFF)
                for info in infos.values()
            ):
                raise ValueError("invalid archive member")
            archive_json = bundle.read("archive.json")
            workbook_bytes = bundle.read("archive.xlsx")
            inner_manifest_bytes = bundle.read("manifest.json")
    except (
        OSError,
        UnicodeError,
        ValueError,
        KeyError,
        RuntimeError,
        zipfile.BadZipFile,
    ) as exc:
        raise PracticeRetentionSafeguardError("练习归档 ZIP 结构无效。") from exc
    if inner_manifest_bytes != outer_manifest_bytes:
        raise PracticeRetentionSafeguardError("练习归档内外 manifest 不一致。")
    try:
        payload = json.loads(archive_json)
        inner_manifest = json.loads(inner_manifest_bytes)
    except (UnicodeError, json.JSONDecodeError) as exc:
        raise PracticeRetentionSafeguardError("练习归档 JSON 无效。") from exc
    if not isinstance(manifest, dict) or manifest != inner_manifest:
        raise PracticeRetentionSafeguardError("练习归档 manifest 无效。")
    if set(manifest) != ARCHIVE_MANIFEST_KEYS:
        raise PracticeRetentionSafeguardError("练习归档 manifest 字段不完整。")
    if (
        manifest.get("schema_version") != 1
        or manifest.get("kind") != "practice-retention"
        or manifest.get("artifact_id") != artifact_id
        or not isinstance(payload, dict)
        or payload.get("schema_version") != 1
        or payload.get("kind") != "practice-retention"
    ):
        raise PracticeRetentionSafeguardError("练习归档 schema 或类型无效。")
    if set(payload) != ARCHIVE_PAYLOAD_KEYS:
        raise PracticeRetentionSafeguardError("练习归档 JSON 字段不完整。")
    files = manifest.get("files")
    if not isinstance(files, dict) or set(files) != {"archive.json", "archive.xlsx"}:
        raise PracticeRetentionSafeguardError("练习归档文件 manifest 不完整。")
    for filename, content in (
        ("archive.json", archive_json),
        ("archive.xlsx", workbook_bytes),
    ):
        digest = files.get(filename)
        if not isinstance(digest, str) or not re.fullmatch(r"[0-9a-f]{64}", digest):
            raise PracticeRetentionSafeguardError("练习归档文件 digest 无效。")
        if hashlib.sha256(content).hexdigest() != digest:
            raise PracticeRetentionSafeguardError("练习归档文件 checksum 校验失败。")

    for key in (
        "candidate_ids",
        "practice_answer_ids",
        "preview_fingerprint",
        "cutoff_at",
        "retention_days",
        "hot_history_limit",
        "hot_history_low_watermark",
    ):
        if payload.get(key) != manifest.get(key):
            raise PracticeRetentionSafeguardError("练习归档内容与 manifest 不一致。")
    candidate_ids = payload.get("candidate_ids")
    answer_ids = payload.get("practice_answer_ids")
    answers = payload.get("practice_answers")
    aggregates = payload.get("aggregate_snapshots")
    candidate_ids_valid = (
        isinstance(candidate_ids, list)
        and bool(candidate_ids)
        and all(
            isinstance(value, int) and not isinstance(value, bool) and value > 0
            for value in candidate_ids
        )
        and candidate_ids == sorted(set(candidate_ids))
    )
    answer_ids_valid = (
        isinstance(answer_ids, list)
        and bool(answer_ids)
        and all(
            isinstance(value, int) and not isinstance(value, bool) and value > 0
            for value in answer_ids
        )
        and len(answer_ids) == len(set(answer_ids))
    )
    if (
        not candidate_ids_valid
        or not answer_ids_valid
        or not isinstance(answers, list)
        or len(answers) != len(answer_ids)
        or not isinstance(aggregates, list)
    ):
        raise PracticeRetentionSafeguardError("练习归档明细内容无效。")
    answer_ids_from_rows = [row.get("id") for row in answers if isinstance(row, dict)]
    if answer_ids_from_rows != answer_ids or len(answer_ids_from_rows) != len(answers):
        raise PracticeRetentionSafeguardError("练习归档明细 ID 不一致。")
    return payload, manifest, archive_sha256


def _validate_backup_after_archive(
    backup_id: str, archive_created_at: datetime
) -> None:
    if BACKUP_ID_PATTERN.fullmatch(backup_id) is None:
        raise PracticeRetentionSafeguardError("配对备份标识无效。")
    backup_path = Path(settings.backup_storage_dir) / backup_id
    try:
        backup_manifest = validate_backup(backup_path)
        backup_created_at = to_utc(
            datetime.fromisoformat(str(backup_manifest["created_at"]))
        )
    except (BackupValidationError, KeyError, TypeError, ValueError) as exc:
        raise PracticeRetentionSafeguardError("配对备份未通过校验。") from exc
    if backup_created_at <= archive_created_at:
        raise PracticeRetentionSafeguardError("配对备份必须创建在练习归档之后。")


def _locked_candidates(db: Session, candidate_ids: list[int]) -> None:
    candidates = (
        db.query(Candidate)
        .filter(Candidate.id.in_(candidate_ids))
        .order_by(Candidate.id)
        .with_for_update()
        .all()
    )
    if [candidate.id for candidate in candidates] != candidate_ids:
        raise PracticeRetentionSafeguardError("受影响考试人不存在，拒绝删除。")


def delete_practice_retention(
    db: Session,
    *,
    candidate_ids: list[int],
    preview_fingerprint: str,
    archive_id: str,
    backup_id: str,
    confirmation: str,
    operator_subject: str,
    now: datetime | None = None,
) -> PracticeRetentionDeleteRead:
    normalized_ids = _normalize_candidate_ids(candidate_ids)
    assert_admin_mutation_allowed(db)
    try:
        payload, manifest, _archive_sha256 = _load_archive(archive_id)
        if (
            manifest.get("candidate_ids") != normalized_ids
            or manifest.get("preview_fingerprint") != preview_fingerprint
            or payload.get("candidate_ids") != normalized_ids
        ):
            raise PracticeRetentionSafeguardError(
                "练习归档与当前所选考试人或预览不匹配。"
            )
        expected_confirmation = (
            f"DELETE PRACTICE ANSWERS {','.join(map(str, normalized_ids))}"
        )
        if confirmation != expected_confirmation:
            raise PracticeRetentionSafeguardError("删除确认文本不匹配。")
        archive_created_at = to_utc(datetime.fromisoformat(str(manifest["created_at"])))
        _validate_backup_after_archive(backup_id, archive_created_at)

        _locked_candidates(db, normalized_ids)
        deleted_at = to_utc(now or datetime.now(UTC))
        preview = preview_practice_retention(
            db,
            now=deleted_at,
            candidate_ids=normalized_ids,
        )
        if preview.fingerprint != preview_fingerprint:
            raise PracticeRetentionSafeguardError("保留预览已过期，请重新预览。")
        answer_ids = _selected_ids_for_candidates(preview, normalized_ids)
        if payload.get("practice_answer_ids") != answer_ids:
            raise PracticeRetentionSafeguardError("练习归档明细与当前预览不匹配。")

        current_answers = (
            db.query(PracticeAnswer)
            .filter(PracticeAnswer.id.in_(answer_ids))
            .order_by(
                PracticeAnswer.candidate_id,
                PracticeAnswer.practiced_at,
                PracticeAnswer.id,
            )
            .all()
        )
        archived_answers = payload["practice_answers"]
        if len(current_answers) != len(archived_answers) or any(
            _answer_snapshot(answer) != archived
            for answer, archived in zip(current_answers, archived_answers, strict=True)
        ):
            raise PracticeRetentionSafeguardError("练习明细内容已变化，拒绝删除。")

        pairs = {
            (answer.candidate_id, answer.question_id) for answer in current_answers
        }
        current_aggregates = (
            db.query(PracticeAnswerAggregate)
            .filter(PracticeAnswerAggregate.candidate_id.in_(normalized_ids))
            .order_by(
                PracticeAnswerAggregate.candidate_id,
                PracticeAnswerAggregate.question_id,
            )
            .all()
        )
        aggregate_by_pair = {
            (aggregate.candidate_id, aggregate.question_id): aggregate
            for aggregate in current_aggregates
            if (aggregate.candidate_id, aggregate.question_id) in pairs
        }
        archived_aggregates = payload["aggregate_snapshots"]
        current_snapshots = [
            _aggregate_snapshot(aggregate_by_pair[pair])
            for pair in sorted(aggregate_by_pair)
        ]
        if set(aggregate_by_pair) != pairs or current_snapshots != archived_aggregates:
            raise PracticeRetentionSafeguardError("练习聚合状态已变化，拒绝删除。")

        deleted_count = (
            db.query(PracticeAnswer)
            .filter(
                PracticeAnswer.id.in_(answer_ids),
                PracticeAnswer.candidate_id.in_(normalized_ids),
            )
            .delete(synchronize_session=False)
        )
        if deleted_count != len(answer_ids):
            raise PracticeRetentionSafeguardError("练习明细删除数量不一致，拒绝提交。")
        record_admin_event(
            db,
            operator_subject=operator_subject,
            action="practice_retention_deleted",
            target_type="practice_answer_set",
            target_id=",".join(map(str, normalized_ids)),
            metadata={
                "archive_ref": archive_id,
                "backup_ref": backup_id,
                "count": deleted_count,
                "deleted_count": deleted_count,
                "fingerprint": preview.fingerprint,
            },
        )
        db.commit()
        db.expire_all()
        return PracticeRetentionDeleteRead(
            deleted_candidate_ids=normalized_ids,
            deleted_practice_answer_ids=answer_ids,
            deleted_count=deleted_count,
            archive_id=archive_id,
            backup_id=backup_id,
        )
    except PracticeRetentionSafeguardError:
        db.rollback()
        raise
    except (BackupValidationError, KeyError, TypeError, ValueError) as exc:
        db.rollback()
        raise PracticeRetentionSafeguardError("练习保留删除校验失败。") from exc
    except Exception:
        db.rollback()
        raise


# Keep names parallel to the existing exam retention service for callers that
# use "retained" terminology.
preview_practice_answer_retention = preview_practice_retention
delete_retained_practice_answers = delete_practice_retention
